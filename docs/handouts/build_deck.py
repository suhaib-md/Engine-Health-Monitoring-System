"""
Builds IgniSense-Deck.pptx (the showcase deck) from the screenshots in ./img.

    pip install python-pptx pillow
    python docs/handouts/build_deck.py [output.pptx]

Edit the SLIDES section at the bottom to change wording; the helpers above it handle layout.
Fonts are Segoe UI / Consolas (stock on Windows). Everything is a native, editable PowerPoint shape.
"""
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"
CROPS = HERE / "img_crops"
CROPS.mkdir(exist_ok=True)
import sys
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "IgniSense-Deck.pptx"

# ---- IgniSense design tokens (src/ui/tokens.ts) ----
BG, PANEL, RAISED, LINE = "0c131c", "121b26", "1a2532", "26323f"
FG, FG2, FG3 = "eef3f8", "b3c0cf", "7f8fa3"
ACCENT, OK, WATCH, WARN, CRIT, VIOLET = "35e0f0", "3ddc84", "f5d547", "ff9a3c", "ff4d5e", "a98bff"
SANS, MONO = "Segoe UI", "Consolas"

W, H = 13.333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
TOTAL = 17  # for the footer counter


def rgb(h):
    return RGBColor.from_string(h.upper())


def new_slide(notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def rect(s, x, y, w, h, fill=PANEL, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.shadow.inherit = False
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid()
        r.fill.fore_color.rgb = rgb(fill)
    if line:
        r.line.color.rgb = rgb(line)
        r.line.width = Pt(lw)
    else:
        r.line.fill.background()
    return r


def text(s, x, y, w, h, runs, size=18, color=FG, bold=False, font=SANS, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, space_after=0):
    """runs: a string, or a list of paragraphs; each paragraph is a string or a list of (text, {overrides})."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        para.space_after = Pt(space_after)
        parts = p if isinstance(p, list) else [(p, {})]
        for t, o in parts:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color))
    return tb


_page = [1]


def chrome(s, kicker):
    """Kicker (top-left, auto-numbered) and footer (brand, page number) on every content slide."""
    _page[0] += 1
    n = _page[0]
    text(s, 0.6, 0.42, 8, 0.3, f"{n - 1:02d} · {kicker}".upper(), size=12, color=ACCENT, bold=True, font=MONO)
    rect(s, 0.6, 7.0, W - 1.2, 0.01, fill=LINE)
    text(s, 0.6, 7.08, 6, 0.25, "IGNISENSE  ·  TEAM REVORA", size=10, color=FG3, font=MONO)
    text(s, W - 2.6, 7.08, 2.0, 0.25, f"{n:02d} / {TOTAL}", size=10, color=FG3, font=MONO, align=PP_ALIGN.RIGHT)


def title(s, t, sub=None, w=12.1, size=30):
    text(s, 0.6, 0.78, w, 1.0, t, size=size, bold=True, color=FG, spacing=1.0)
    if sub:
        text(s, 0.6, 1.62, w, 0.6, sub, size=17, color=FG2)


_FONTS = {}


def _font(size, bold=False):
    key = (round(size), bold)
    if key not in _FONTS:
        try:
            _FONTS[key] = ImageFont.truetype("segoeuib.ttf" if bold else "segoeui.ttf", round(size * 4))
        except OSError:
            _FONTS[key] = None
    return _FONTS[key]


def measure_lines(t, size, width_in, bold=False):
    """Number of wrapped lines for text t at `size` pt in a box `width_in` wide (Segoe UI metrics when available)."""
    f = _font(size, bold)
    max_w = width_in * 72 * 4  # font was loaded at 4x
    words, lines, cur = t.split(" "), 1, ""
    for w_ in words:
        trial = (cur + " " + w_).strip()
        width = f.getlength(trial) if f else len(trial) * size * 0.55 * 4
        if width > max_w and cur:
            lines += 1
            cur = w_
        else:
            cur = trial
    return lines


def bullets(s, x, y, w, items, size=18, gap=12, color=FG2, mark=ACCENT):
    """Square-marker bullets, laid out by measured text height. items: strings or [(text, overrides)...]."""
    cy = y
    for it in items:
        plain = it if isinstance(it, str) else "".join(t for t, _ in it)
        lines = measure_lines(plain, size, w - 0.3 - 0.05)
        h = lines * size * 1.05 * 1.2 / 72
        rect(s, x, cy + 0.09, 0.11, 0.11, fill=mark)
        text(s, x + 0.3, cy, w - 0.3, h, [it], size=size, color=color, spacing=1.05)
        cy += h + gap / 72
    return cy


def crop(name, box=None):
    """Return the path of an image, optionally cropped to (left, top, right, bottom) fractions."""
    src = IMG / name
    if not box:
        return src
    dst = CROPS / (src.stem + "_" + "_".join(f"{int(b * 100)}" for b in box) + ".png")
    im = Image.open(src)
    w, h = im.size
    im.crop((int(box[0] * w), int(box[1] * h), int(box[2] * w), int(box[3] * h))).save(dst)
    return dst


def picture(s, name, x, y, w, h, box=None, border=LINE):
    """Place an image inside the (x, y, w, h) frame, keeping its aspect ratio, with a thin border."""
    path = crop(name, box)
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    rect(s, px - 0.03, py - 0.03, pw + 0.06, ph + 0.06, fill=border)
    s.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    return px, py, pw, ph


def badge(s, x, y, n, color=ACCENT):
    c = rect(s, x, y, 0.36, 0.36, fill=color, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size, r.font.bold, r.font.name = Pt(14), True, MONO
    r.font.color.rgb = rgb("06222a")
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def stat(s, x, y, w, big, label, color=ACCENT, size=54):
    text(s, x, y, w, 0.9, big, size=size, bold=True, color=color, font=MONO)
    text(s, x, y + 0.95, w, 0.8, label, size=15, color=FG2)


def logo(s, x, y, size=0.7):
    k = size / 40
    pts = [(14, 8), (30, 8), (26, 32), (10, 32)]
    fb = s.shapes.build_freeform(Inches(x + pts[0][0] * k), Inches(y + pts[0][1] * k))
    fb.add_line_segments([(Inches(x + px * k), Inches(y + py * k)) for px, py in pts[1:]], close=True)
    shp = fb.convert_to_shape()
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(ACCENT)
    shp.line.fill.background()
    rect(s, x + 27 * k, y + 27 * k, 6 * k, 6 * k, fill=CRIT)


def node(s, x, y, w, h, name, sub, fill=PANEL, line=LINE, name_color=ACCENT, sub_color=FG2):
    rect(s, x, y, w, h, fill=fill, line=line, lw=1.25)
    text(s, x + 0.15, y + 0.14, w - 0.3, 0.4, name, size=17, bold=True, color=name_color, font=MONO, align=PP_ALIGN.CENTER)
    text(s, x + 0.15, y + 0.6, w - 0.3, h - 0.65, sub, size=13, color=sub_color, align=PP_ALIGN.CENTER, spacing=1.05)


def arrow(s, x, y, w=0.45, color=FG3):
    a = rect(s, x, y, w, 0.3, fill=color, shape=MSO_SHAPE.RIGHT_ARROW)
    return a


# =====================================================================================
#                                       SLIDES
# =====================================================================================

# 1 ---- title
s = new_slide("Welcome. IgniSense is a digital twin of a petrol engine that finds faults from sensor data alone. "
              "Everything you'll see runs live in this browser, with no server behind it.")
rect(s, 0, 0, 0.18, H, fill=ACCENT)
logo(s, 0.9, 0.85, 0.9)
text(s, 1.95, 0.98, 6, 0.5, "IGNISENSE", size=22, bold=True, color=ACCENT, font=MONO)
text(s, 0.9, 2.2, 11, 1.6, ["Smart engine", "health diagnostic"], size=60, bold=True, color=FG, spacing=0.95)
rect(s, 0.9, 4.2, 1.6, 0.06, fill=ACCENT)
text(s, 0.9, 4.5, 9.2, 1.0, "A digital twin that finds engine faults, names the failing part, and says how long is left, "
     "using nothing but sensor readings.", size=21, color=FG2, spacing=1.15)
text(s, 0.9, 6.35, 6, 0.4, "TEAM REVORA", size=16, bold=True, color=FG, font=MONO)
text(s, 0.9, 6.7, 9, 0.35, "Physics-informed engine digital twin  ·  runs entirely in the browser", size=12, color=FG3, font=MONO)

# 2 ---- problem
s = new_slide("Service on a fixed schedule either replaces parts that are still good or misses a part that fails early. "
              "Condition monitoring watches the engine itself. Our project shows how a monitor can do that from sensors alone.")
chrome(s, "The problem")
title(s, "Engines fail. Schedules don't know when.")
rect(s, 0.6, 2.2, 5.85, 3.3, fill=PANEL, line=LINE)
text(s, 0.95, 2.45, 5.2, 0.4, "FIXED-INTERVAL SERVICE", size=13, bold=True, color=WARN, font=MONO)
text(s, 0.95, 2.95, 5.2, 2.4, ["Changes parts that are still good.", "", "Misses the part that fails early.", "",
                                "Says nothing about what is wrong, only that it is time."], size=19, color=FG2, spacing=1.05)
rect(s, 6.88, 2.2, 5.85, 3.3, fill=PANEL, line=ACCENT, lw=1.5)
text(s, 7.23, 2.45, 5.2, 0.4, "CONDITION MONITORING", size=13, bold=True, color=ACCENT, font=MONO)
text(s, 7.23, 2.95, 5.2, 2.4, ["Watches the engine's own sensors.", "", "Notices a fault while it is small.", "",
                                "Names the part, and estimates the time left."], size=19, color=FG, spacing=1.05)
text(s, 0.6, 5.95, 12.1, 0.8, [[("IgniSense shows how, ", {}), ("from sensor readings alone.", {"color": ACCENT, "bold": True})]],
     size=24, color=FG)

# 3 ---- idea / architecture
s = new_slide("Two engines run side by side. The Plant is the real engine and holds the hidden fault. The Twin is a healthy copy of the same physics that only sees the sensors. "
              "The gap between what the sensors read and what the twin expects is the evidence. The analytics can't see the fault, and a lint rule enforces that.")
chrome(s, "The idea")
title(s, "Two engines. One rule.")
node(s, 0.6, 2.2, 2.35, 1.55, "PLANT", "the \"real\" engine\nhidden faults live here", fill=PANEL, line=CRIT, name_color=CRIT)
arrow(s, 3.02, 2.85)
node(s, 3.55, 2.2, 2.35, 1.55, "SENSORS", "noise, drift,\nspikes, dropouts")
arrow(s, 5.97, 2.85)
node(s, 6.5, 2.2, 2.35, 1.55, "TELEMETRY", "one common format\nfor every source", fill="10333d", line=ACCENT)
arrow(s, 8.92, 2.85)
node(s, 9.45, 1.95, 3.25, 1.0, "TWIN", "what the sensors\nshould read", name_color=FG)
node(s, 9.45, 3.1, 3.25, 1.0, "ANALYTICS", "gap → fault, health, time left", name_color=ACCENT)
rect(s, 0.6, 4.6, 12.1, 1.75, fill=PANEL, line=OK, lw=1.5)
text(s, 0.95, 4.82, 11.4, 0.4, "THE RULE THAT MAKES IT HONEST", size=13, bold=True, color=OK, font=MONO)
text(s, 0.95, 5.3, 11.4, 1.0, "The analytics only ever receive sensor telemetry. The fault lives in a part of the code they are not allowed "
     "to import, and a lint rule enforces it. The system cannot cheat by reading the answer.", size=19, color=FG, spacing=1.1)
text(s, 0.6, 6.5, 12.1, 0.4, "The gap between what the sensors read and what a healthy engine would read is the evidence.",
     size=15, color=FG3)

# 4 ---- live twin
s = new_slide("This is the live twin page. The 3D engine is procedural and moved by the same slider-crank equation as the numbers; it is slowed 100 times so you can see it. "
              "On the right the health ring and the six subsystem scores.")
chrome(s, "See it live")
title(s, "A 3D engine on the same physics", w=4.7)
bullets(s, 0.6, 2.75, 4.7, [
    "Pistons, rods and crank move by the slider-crank equation, not an animation loop.",
    "Cylinders flash in firing order 1-3-4-2.",
    "Health ring and six subsystem scores update live.",
    "Click any part for its sensors and its calculation.",
], size=17)
picture(s, "live-running.png", 5.65, 1.05, 7.15, 5.75)

# 5 ---- gauges
s = new_slide("Every gauge carries a Twin marker: the white tick is what a healthy engine would read right now, at the same speed, load and ambient. "
              "The delta turns coloured once it leaves the noise. The tachometer is a car-style dial. Click any number and you get the equation.")
chrome(s, "The signals")
title(s, "Every gauge has a reference")
bullets(s, 0.6, 1.95, 3.9, [
    [("White tick", {"color": FG, "bold": True}), (" = what a healthy engine would read right now.", {})],
    [("Δ", {"color": FG, "bold": True, "font": MONO}), (" turns coloured once the gap leaves the noise band.", {})],
    "A dead sensor shows — —, never zero.",
    [("Click any value", {"color": FG, "bold": True}), (" to see the equation behind it.", {})],
], size=17)
picture(s, "live-signals.png", 4.75, 1.85, 8.05, 4.95)

# 5b ---- break it
s = new_slide("Inject any fault from the test bench. The sump and oil gallery pulse red, the ring drops and turns orange, and the card names the fault. "
              "Nothing on screen told the monitor which fault we chose. It found it in the sensors.")
chrome(s, "Break it, and it shows")
title(s, "Inject a fault. Watch it show.", w=4.7)
bullets(s, 0.6, 2.75, 4.7, [
    "Oil-pump wear injected from the test bench.",
    "The sump and oil gallery pulse red; the health ring drops and turns orange.",
    [("The monitor names ", {}), ("Lubrication-system degradation", {"color": FG, "bold": True}), (" from the sensors alone.", {})],
    "Repair, and everything settles back.",
], size=17)
picture(s, "live-fault.png", 5.65, 1.05, 7.15, 5.75)

# 6 ---- diagnosis
s = new_slide("When we inject an oil-pump fault, the monitor names it and lists the evidence. Notice we call it an evidence score, never a confidence percentage, "
              "because it is a weighted sum, not a calibrated probability.")
chrome(s, "Diagnosis")
title(s, "It names the fault, and shows its evidence")
bullets(s, 0.6, 1.95, 3.9, [
    "Probable fault, severity, and a diagnostic evidence score.",
    "\"Why the system thinks this\": named symptoms, each colour-coded.",
    "A recommended action and the model's own confidence.",
    "Alerts escalate WATCH → WARNING → CRITICAL, with hysteresis.",
], size=17)
picture(s, "live-diagnosis.png", 4.75, 1.85, 8.05, 4.95)

# 7 ---- early warning
s = new_slide("CUSUM adds up small persistent shifts. A residual that sits at one sigma never crosses a three-sigma limit, but CUSUM passes its limit in about eleven samples. "
              "So the monitor warns before any fixed threshold trips.")
chrome(s, "Early warning")
title(s, "It warns before any limit trips")
bullets(s, 0.6, 1.95, 3.9, [
    [("CUSUM", {"color": FG, "bold": True}), (" adds up small, persistent shifts.", {})],
    "A gap of 1σ never crosses a 3σ limit. CUSUM crosses its limit in about 11 samples.",
    [("D²", {"color": FG, "bold": True}), (" flags an unusual combination of residuals (χ² limit 15.09).", {})],
    "Violet lines mark injections. Diamonds mark alerts.",
], size=17)
picture(s, "trends-fault.png", 4.75, 1.85, 8.05, 4.95)

# 8 ---- misfire
s = new_slide("A misfire removes one power pulse every two revolutions. That repeats at half the crank frequency. The cylinders fire 180 degrees of crank apart, "
              "which is 90 degrees of phase at half order, so the phase of that half-order line points at the cylinder. And the torque command rises by exactly four thirds.")
chrome(s, "Which cylinder?")
title(s, "One missing pulse points at its cylinder")
bullets(s, 0.6, 1.95, 3.9, [
    "A misfire repeats every two revolutions: a 0.5× line.",
    "That line's phase names the cylinder.",
    [("Torque command rises by exactly ", {}), ("4/3", {"color": ACCENT, "bold": True}), (": three cylinders do four's work.", {})],
], size=17)
rect(s, 0.6, 4.7, 3.9, 2.05, fill=PANEL, line=LINE)
text(s, 0.85, 4.82, 3.4, 0.3, "0.5× PHASE  →  CYLINDER", size=11, bold=True, color=ACCENT, font=MONO)
for i, (c, ph) in enumerate([("C1", "+45°"), ("C2", "+135°"), ("C3", "−45°"), ("C4", "−135°")]):
    yy = 5.2 + i * 0.36
    text(s, 0.85, yy, 1.4, 0.3, c, size=15, bold=True, color=FG, font=MONO)
    text(s, 2.4, yy, 1.8, 0.3, ph, size=15, color=WARN if c == "C3" else FG2, font=MONO, bold=(c == "C3"))
picture(s, "vibration-misfire.png", 4.75, 1.85, 8.05, 4.95)

# 9 ---- sensor lies
s = new_slide("If the coolant reading jumps 25 degrees in a tenth of a second, either the engine did something impossible or the sensor failed. "
              "The fastest the coolant can heat is all the coolant heat at full load with no cooling: 81.5 kilowatts into 100 kilojoules per kelvin, so 0.8 K per second. "
              "The reading beat that by over 300 times. So we blame the sensor, not the engine.")
chrome(s, "Sensor faults")
title(s, "When a sensor lies, physics catches it")
stat(s, 0.6, 2.3, 3.8, "0.8 K/s", "the fastest coolant can possibly heat: all coolant heat at full load, no cooling", color=OK)
stat(s, 4.75, 2.3, 3.8, "250 K/s", "what the reading implied: +25 °C in 0.1 s", color=CRIT)
stat(s, 8.9, 2.3, 3.8, "300×", "over the physical limit, so the sensor failed, not the engine", color=WARN)
text(s, 0.6, 4.05, 12.1, 0.5, "dT/dt|max = Q_cool,max / C_th = 81 500 W / 100 000 J/K ≈ 0.8 K/s", size=20, color=FG2, font=MONO)
rect(s, 0.6, 4.95, 12.1, 1.55, fill=PANEL, line=ACCENT, lw=1.5)
text(s, 0.95, 5.13, 11.4, 0.35, "THE VERDICT ON SCREEN", size=12, bold=True, color=ACCENT, font=MONO)
text(s, 0.95, 5.55, 11.4, 0.9, [[("Coolant sensor fault", {"bold": True, "color": FG}),
                                   (". The cooling system is not blamed: the other sensors and the Twin agree with each other, "
                                    "and the bad readings are left out of the diagnosis.", {})]], size=19, color=FG2, spacing=1.1)

# 10 ---- calculations
s = new_slide("Every number has an equation behind it, and you can click it. The drawer shows the chain with today's live numbers, and a check table that proves the result equals the gauge. "
              "We test that in six different engine states.")
chrome(s, "Calculations")
title(s, "Click any number. See the equation.")
bullets(s, 0.6, 1.95, 3.9, [
    "28 registered equations, each with symbols and units.",
    "Live numbers substituted into the LaTeX.",
    [("A check table proves it: ", {}), ("✓ EQUAL", {"color": OK, "bold": True, "font": MONO}), (" to the gauge.", {})],
    "Tested in six engine states.",
], size=17)
picture(s, "math-drawer.png", 4.75, 1.85, 8.05, 4.95)

# 11 ---- blind test
s = new_slide("This is the evidence that it isn't reading the injected fault. An audience member selects a sealed case; the fault is chosen inside the simulation worker, and the interface never learns what is in it. "
              "The monitor has to identify it, and then the answer is revealed.")
chrome(s, "Blind test")
title(s, "Test it blind.")
bullets(s, 0.6, 1.95, 3.9, [
    "Nine sealed fault cases, shuffled inside the simulation.",
    "The interface never sees what is in the case.",
    "The monitor names the fault from sensors alone.",
    [("Reveal → ", {}), ("✓ CORRECT", {"color": OK, "bold": True, "font": MONO}), (" or not.", {})],
], size=17)
picture(s, "blind.png", 4.75, 1.85, 8.05, 4.95)

# 12 ---- RUL + report
s = new_slide("The remaining-life estimate is a straight-line trend with a plus or minus two standard error band, shown only when the slope is statistically real. "
              "Raise the load and the wear speeds up, so the remaining life shrinks. The one-page report prints to PDF.")
chrome(s, "How long is left?")
title(s, "Remaining life, with an honest band")
bullets(s, 0.6, 1.95, 3.9, [
    "A trend line with a ±2 standard-error band.",
    [("Shown only ", {}), ("when the trend is significant", {"color": FG, "bold": True}), (".", {})],
    "Higher load wears faster, so the life shrinks.",
    "One-click, one-page maintenance report as a PDF.",
], size=17)
picture(s, "report.png", 4.75, 1.85, 8.05, 4.95)

# 13 ---- proof
s = new_slide("We don't just claim the physics is right. The Validation page recomputes hand-calculated values live in the browser, and every one passes. "
              "There are also 363 automated tests, and every golden number from the research review reproduces within two percent.")
chrome(s, "Checked, not just claimed")
title(s, "Every number is checked against hand calculation")
stat(s, 0.6, 2.0, 3.0, "363", "automated tests, all passing", color=OK, size=48)
stat(s, 0.6, 3.75, 3.0, "26", "live checks on the Validation page", color=ACCENT, size=48)
stat(s, 0.6, 5.5, 3.0, "2%", "tolerance on every review value", color=WATCH, size=48)
picture(s, "validation.png", 4.2, 1.85, 8.6, 4.95)

# 14 ---- hardware
s = new_slide("Every source emits the same telemetry object, so adding real hardware means adding one adapter file. "
              "Replay of a recorded CSV is tested end to end and reproduces the diagnosis. ESP32 and OBD-II are tested against parsers and simulated streams, not yet real hardware.")
chrome(s, "From simulator to a real engine")
title(s, "Same analytics, any source")
srcs = [("SIMULATOR", "built-in plant", OK), ("CSV REPLAY", "recorded runs", OK), ("ESP32", "USB, real sensors", WATCH),
        ("OBD-II", "a real car", WATCH), ("PHONE", "accelerometer", WATCH)]
for i, (n, sub, col) in enumerate(srcs):
    yy = 2.0 + i * 0.92
    rect(s, 0.6, yy, 3.0, 0.75, fill=PANEL, line=col, lw=1.25)
    text(s, 0.8, yy + 0.08, 2.6, 0.3, n, size=14, bold=True, color=FG, font=MONO)
    text(s, 0.8, yy + 0.4, 2.6, 0.3, sub, size=12, color=FG2)
    arrow(s, 3.75, yy + 0.22, 0.6, FG3)
node(s, 4.55, 2.6, 3.0, 1.8, "TELEMETRY", "one common\nformat", fill="10333d", line=ACCENT)
arrow(s, 7.7, 3.35, 0.6, FG3)
node(s, 8.5, 2.2, 4.2, 1.15, "TWIN", "healthy reference", name_color=FG)
node(s, 8.5, 3.6, 4.2, 1.15, "ANALYTICS", "unchanged for every source", name_color=ACCENT)
text(s, 4.55, 5.15, 8.15, 1.4, [[("Green", {"color": OK, "bold": True}), (" = tested end to end.   ", {}),
                                  ("Yellow", {"color": WATCH, "bold": True}), (" = built and tested on parsers and simulated streams; real hardware is the next step.", {})]],
     size=15, color=FG2, spacing=1.15)

# 15 ---- limits and next
s = new_slide("We would rather state the limits ourselves. The values are demo calibration, and the data is synthetic. The next step is a heat-balance run on a real engine rig to replace demo calibration with measured data.")
chrome(s, "Honest limits, clear next steps")
title(s, "What it is not, and what comes next")
rect(s, 0.6, 2.0, 5.85, 4.55, fill=PANEL, line=LINE)
text(s, 0.95, 2.22, 5.2, 0.35, "LIMITS, STATED PLAINLY", size=13, bold=True, color=WARN, font=MONO)
bullets(s, 0.95, 2.75, 5.2, [
    "Synthetic telemetry; every value is demo calibration.",
    "Lumped thermal model: one coolant, one oil temperature.",
    "Evidence scores are weighted sums, not probabilities.",
    "Remaining life assumes the load continues.",
], size=16, gap=8, mark=WARN)
rect(s, 6.88, 2.0, 5.85, 4.55, fill=PANEL, line=ACCENT, lw=1.5)
text(s, 7.23, 2.22, 5.2, 0.35, "NEXT", size=13, bold=True, color=ACCENT, font=MONO)
bullets(s, 7.23, 2.75, 5.2, [
    "Heat-balance run on the college engine rig, to replace demo calibration with measured values.",
    "ESP32 with DS18B20, MAX31855, a pressure transducer and an ADXL345.",
    "Log real failure data, then evaluate machine learning against our statistics.",
], size=16, gap=8)

# 16 ---- close
s = new_slide("Invite the audience to choose a fault for you to inject, or to select a blind-test case.")
rect(s, 0, 0, 0.18, H, fill=ACCENT)
logo(s, 0.9, 0.85, 0.9)
text(s, 1.95, 0.98, 6, 0.5, "IGNISENSE", size=22, bold=True, color=ACCENT, font=MONO)
text(s, 0.9, 2.4, 11.5, 2.4, ["Choose a fault.", [("See it diagnosed.", {"color": ACCENT})]], size=72, bold=True, color=FG, spacing=0.95)
text(s, 0.9, 5.2, 11, 0.5, "Ask us to inject any fault: the monitor only ever sees the sensors.", size=21, color=FG2)
text(s, 0.9, 6.35, 6, 0.4, "TEAM REVORA", size=16, bold=True, color=FG, font=MONO)
text(s, 0.9, 6.7, 9, 0.35, "IgniSense — Smart Engine Health Diagnostic  ·  all values are demo calibration", size=12, color=FG3, font=MONO)

prs.save(OUT)
print("saved", OUT, len(prs.slides._sldIdLst), "slides")
